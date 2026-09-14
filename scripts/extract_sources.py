import io,json,zipfile,datetime,pathlib,hashlib
import openpyxl
ROOT=pathlib.Path(__file__).resolve().parents[1]
SRC=pathlib.Path('C:/Users/jihoo/OneDrive/바탕 화면')
OUT=ROOT/'data'
OUT.mkdir(exist_ok=True)
def val(v):
    if isinstance(v,(datetime.datetime,datetime.date,datetime.time)): return v.isoformat()
    return v
def sheet(s):
    rows=[]
    for i,r in enumerate(s.values,1):
        cells={openpyxl.utils.get_column_letter(j):val(v) for j,v in enumerate(r,1) if v is not None}
        if cells: rows.append({'row':i,'cells':cells})
    return {'name':s.title,'rows':rows}
w=openpyxl.load_workbook(io.BytesIO((SRC/'행사 시트 자동화').read_bytes()),read_only=True,data_only=True)
base=[sheet(s) for s in w]
(OUT/'base.json').write_text(json.dumps(base,ensure_ascii=False),encoding='utf-8')
print('BASE',[(s['name'],len(s['rows'])) for s in base])
for s in base:
    if 'DB' in s['name']: print(s['name'],s['rows'][:2])
    if s['name']=='시트_260831': print('TEMPLATE',s['rows'])
z=zipfile.ZipFile(SRC/'행사시트.zip')
index=[]
for name in z.namelist():
    if not name.endswith('.xlsx'): continue
    wb=openpyxl.load_workbook(io.BytesIO(z.read(name)),read_only=True,data_only=True)
    entries=[]
    for s in wb:
        data=sheet(s)
        ident=hashlib.sha256((name+'|'+s.title).encode()).hexdigest()[:16]
        (OUT/(ident+'.json')).write_text(json.dumps(data,ensure_ascii=False),encoding='utf-8')
        errors=sum(isinstance(v,str) and v.startswith(('#REF!','#VALUE!','#DIV/0!','#N/A','#NAME?','#NUM!')) for r in data['rows'] for v in r['cells'].values())
        entries.append({'id':ident,'name':s.title,'rows':len(data['rows']),'errors':errors})
    index.append({'file':pathlib.PurePosixPath(name).name,'sheets':entries})
    print('ARCHIVE',index[-1]['file'],len(entries),flush=True)
(OUT/'archive.json').write_text(json.dumps(index,ensure_ascii=False),encoding='utf-8')
