"""Read the supplied XLSX and ZIP into typed, traceable website data (never edit sources)."""
import datetime, hashlib, io, json, pathlib, zipfile
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
WORKBOOK = pathlib.Path('C:/Users/jihoo/Downloads/26년 행사시트 디벨롭 (1).xlsx')
ZIP = pathlib.Path('C:/Users/jihoo/OneDrive/바탕 화면/행사시트.zip')
def value(v):
    return v.isoformat() if isinstance(v, (datetime.datetime, datetime.date, datetime.time)) else v
def number(v):
    return v if isinstance(v, (float,int)) and not isinstance(v,bool) else None
def label(v):
    return str(v or '').replace(' ','').replace('\n','')
def rows(s):
    return [{'row':i,'cells':{openpyxl.utils.get_column_letter(j):value(v) for j,v in enumerate(r,1) if v is not None}} for i,r in enumerate(s.values,1) if any(v is not None for v in r)]

cached=openpyxl.load_workbook(WORKBOOK, data_only=True)
formulas=openpyxl.load_workbook(WORKBOOK, data_only=False)
base=json.loads((ROOT/'data/base.json').read_text(encoding='utf-8'))
for source in base:
    if source['name'] in cached.sheetnames: source['rows']=rows(cached[source['name']])
sheet=next(s for s in formulas if '0831' in s.title)
values=cached[sheet.title]
mapping={'workbook':WORKBOOK.name,'sheet':sheet.title,'sha256':hashlib.sha256(WORKBOOK.read_bytes()).hexdigest(),
 'profitRows':[{'row':r,'label':sheet.cell(r,7).value,'target':sheet.cell(r,8).value,'actual':sheet.cell(r,9).value,'delta':sheet.cell(r,10).value,'achievement':sheet.cell(r,11).value} for r in range(18,36)],
 'salesColumns':[{'column':openpyxl.utils.get_column_letter(c),'label':sheet.cell(4,c).value,'formula':sheet.cell(6,c).value if sheet.cell(6,c).data_type=='f' else None} for c in range(15,36)],
 'defaults':{'invoiceCount':number(values['K15'].value),'targetAdRate':number(values['H32'].value),'actualAdRate':number(values['I32'].value)}}
(ROOT/'data/base.json').write_text(json.dumps(base,ensure_ascii=False),encoding='utf-8')
(ROOT/'data/workbook-map.json').write_text(json.dumps(mapping,ensure_ascii=False,indent=2),encoding='utf-8')

# The archive's original summary/media cells take precedence over flattened legacy DB metrics.
archive=zipfile.ZipFile(ZIP)
archive_books={}
events=[]
perf=next(s for s in base if 'PERFORMANCE_DB' in s['name'])['rows'][1:]
media_db=next(s for s in base if 'MEDIA_PLAN_DB' in s['name'])['rows'][1:]
seen=set()
for record in perf:
    c=record['cells']; origin=str(c.get('R','')); parts=origin.split(' | ',1)
    if len(parts)!=2: continue
    filename,title=parts
    key=(filename,title)
    if key in seen: continue
    seen.add(key)
    member=next((n for n in archive.namelist() if n.endswith(filename)),None)
    if not member: continue
    if member not in archive_books: archive_books[member]=openpyxl.load_workbook(io.BytesIO(archive.read(member)),data_only=True)
    book=archive_books[member]
    if title not in book.sheetnames: continue
    s=book[title]; summary={}; media=[]; conversion=None
    summary_start=next((r for r in range(1,min(s.max_row,180)+1) if '총 판매' in str(s.cell(r,2).value)),None)
    if summary_start:
        for r in range(summary_start+1,min(summary_start+19,s.max_row)+1):
            summary_label=str(s.cell(r,2).value or '').replace(' ',''); summary[summary_label]=number(s.cell(r,4).value)
    header=next((r for r in range(1,min(s.max_row,180)+1) if '광고비 상세' in str(s.cell(r,2).value)),None)
    if header:
        for r in range(header+1,min(header+16,s.max_row)+1):
            name=s.cell(r,2).value
            if not isinstance(name,str) or any(x in name for x in ['합계','총계','TOTAL','전환','목표','트래픽','■']): continue
            budget=number(s.cell(r,3).value); target=number(s.cell(r,6).value)
            spent=number(s.cell(r,11).value) if '광고비' in str(s.cell(header,11).value) else None
            traffic=number(s.cell(r,12).value) if '트래픽' in str(s.cell(header,12).value) else None
            if budget is None and spent is None: continue
            media.append({'name':name,'budget':budget,'targetTraffic':target,'spent':spent,'actualTraffic':traffic,'source':f'{filename} | {title}!B{r}:M{r}'})
    # Later event templates put profit and media blocks on the right of the sales list.
    target_metrics={}; actual_metrics={}; data_issues=[]
    for hr in range(1,min(s.max_row,25)+1):
        for gc in range(2,min(s.max_column,80)):
            if label(s.cell(hr,gc).value) not in ['목표','예상']: continue
            if label(s.cell(hr,gc+1).value) not in ['달성','실행','실적']: continue
            target_candidate={}; actual_candidate={}
            for rr in range(hr+1,min(hr+32,s.max_row)+1):
                name=next((label(s.cell(rr,cc).value) for cc in [gc-1,gc-2] if cc>0 and isinstance(s.cell(rr,cc).value,str) and not str(s.cell(rr,cc).value).startswith('#')),None)
                if name:
                    target_candidate[name]=number(s.cell(rr,gc).value)
                    actual_candidate[name]=number(s.cell(rr,gc+1).value)
            if any(k.startswith('공헌이익') for k in actual_candidate):
                target_metrics=target_candidate;actual_metrics=actual_candidate;break
        if actual_metrics:break
    for hr in range(1,min(s.max_row,150)+1):
        columns={label(s.cell(hr,cc).value):cc for cc in range(1,min(s.max_column,80)+1) if isinstance(s.cell(hr,cc).value,str)}
        if '매체' not in columns or '목표트래픽' not in columns:continue
        lookup=lambda keys:next((columns[k] for k in keys if k in columns),None)
        nc=columns['매체'];bc=lookup(['비용','계획비용','예산']);sc=lookup(['실소진','실행비용','실제소진']);tc=columns['목표트래픽'];ac=lookup(['달성트래픽','실제트래픽'])
        typed=[]
        for rr in range(hr+1,min(hr+21,s.max_row)+1):
            name=s.cell(rr,nc).value
            if label(name).lower() in ['total','합계','총계']:break
            if not isinstance(name,str) or not name.strip():continue
            get=lambda cc:number(s.cell(rr,cc).value) if cc else None
            if all(get(cc) is None for cc in [bc,sc,tc,ac]):continue
            typed.append({'name':name.strip(),'budget':get(bc),'targetTraffic':get(tc),'spent':get(sc),'actualTraffic':get(ac),'source':f'{filename} | {title}!{openpyxl.utils.get_column_letter(nc)}{rr}'})
        if typed:media=typed;break
    if not media:
        for m in media_db:
            d=m['cells']
            if d.get('B')==c.get('B'):
                media.append({'name':d.get('H'),'budget':number(d.get('I')),'targetTraffic':number(d.get('K')),'spent':None,'actualTraffic':None,'source':d.get('M')})
    event_title=c.get('D')
    if '행사' in label(s['B2'].value) and ('타이틀' in label(s['B2'].value) or '명분' in label(s['B2'].value)) and isinstance(s['C2'].value,str):event_title=s['C2'].value
    if not event_title or str(event_title).strip() in ['시작일','종료일','행사명','행사 명','채널']: event_title=title
    metric=lambda names:next((actual_metrics[k] for k in names if actual_metrics.get(k) is not None),None)
    target_metric=lambda names:next((target_metrics[k] for k in names if target_metrics.get(k) is not None),None)
    revenue=metric(['고객결제가','고객결제금액','소비자결제가','매출','표면매출'])
    target_revenue=target_metric(['고객결제가','고객결제금액','소비자결제가','매출','표면매출'])
    qty=metric(['판매수량(결제건수)','판매수량','주문수량'])
    profit=metric(['공헌이익']);margin=metric(['공헌이익(%)','공헌이익율','공헌이익률'])
    if revenue is None:revenue=summary.get('매출',number(c.get('H')))
    if qty is None:qty=summary.get('주문수량',number(c.get('I')))
    if profit is None:profit=summary.get('공헌이익')
    if not actual_metrics and not summary:data_issues.append('성과 요약 표의 위치를 확인해야 합니다. 일부 값은 기존 이관 자료입니다.')
    spent=sum(m['spent'] or 0 for m in media)
    roas=revenue/spent if revenue is not None and spent>0 else None
    events.append({'id':c.get('B'),'channel':c.get('C'),'title':event_title,'start':c.get('E'),'end':c.get('F'),
      'targetRevenue':target_revenue,'revenue':revenue,'qty':qty,'profit':profit,'margin':margin,'roas':roas,
      'conversion':conversion,'source':origin,'media':media,'dataIssues':data_issues})
(ROOT/'data/event-history.json').write_text(json.dumps({'source':ZIP.name,'sha256':hashlib.sha256(ZIP.read_bytes()).hexdigest(),'events':events},ensure_ascii=False),encoding='utf-8')
print(json.dumps({'sheets':len(base),'profitRows':len(mapping['profitRows']),'events':len(events),'mediaRows':sum(len(e['media']) for e in events)},ensure_ascii=True))
